"""Create the report companion, combine PDFs, and archive diagram deliverables."""
from __future__ import annotations

import ast
from datetime import datetime
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from pypdf import PdfReader, PdfWriter
from PIL import Image

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent.parent
content = json.loads((HERE / 'nguon/noi-dung.json').read_text())
figures = json.loads((HERE / 'danh-muc-hinh.json').read_text())


def source_link(ref):
    name, line = ref.split(':')
    return f'[{ref}](../../{name}#L{line})'


def md_table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |', *['| ' + ' | '.join(str(c).replace('|', '\\|') for c in row) + ' |' for row in rows]])


class_names = None
for stmt in ast.parse((PROJECT / 'inference.py').read_text()).body:
    if isinstance(stmt, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'CLASS_NAMES' for t in stmt.targets):
        class_names = ast.literal_eval(stmt.value)
assert class_names is not None and len(class_names) == 19
vi_names = ['Nền', 'Đồng cỏ', 'Lúa mì mềm vụ đông', 'Ngô', 'Đại mạch vụ đông', 'Cải dầu vụ đông', 'Đại mạch vụ xuân', 'Hướng dương', 'Nho', 'Củ cải đường', 'Triticale vụ đông', 'Lúa mì cứng vụ đông', 'Trái cây, rau và hoa', 'Khoai tây', 'Cây họ đậu làm thức ăn chăn nuôi', 'Đậu tương', 'Vườn cây ăn quả', 'Ngũ cốc hỗn hợp', 'Cao lương']
class_rows = [[i, name, vi_names[i]] for i, name in enumerate(class_names)]

md = [f'# {content["title"]}', '', content['scope'], '',
      'Ngày đối chiếu: **2026-10-30**. Bộ tài liệu gồm 12 hình; nhánh GEODES là hình phụ lục. Sơ đồ được dựng từ mã nguồn, không lấy luồng cũ mô tả prediction có sẵn trong một số README.', '',
      '## Cách sử dụng trong báo cáo', '',
      '- Dùng **SVG** để giữ nét khi phóng lớn hoặc chèn vào Word có hỗ trợ SVG; dùng **PNG** nếu phần mềm soạn thảo không hiển thị SVG. Các tệp PDF riêng cũng là vector.',
      '- **So-do-Crop-Map.docx** chứa hình, chú thích, thuyết minh, thuật ngữ và bảng API. **So-do-Crop-Map.pdf** là bộ hình để xem/in; mỗi hình nằm trên một trang A3 hoặc A2 tùy mật độ nội dung.',
      '- Tệp **index.html** mở được offline để xem toàn bộ hình và tải từng định dạng. Ảnh lấy từ thư mục hinh; giữ nguyên cấu trúc khi di chuyển tài liệu.',
      '- Nguồn **.puml** dùng PlantUML cho ca sử dụng và hoạt động UML. Nguồn **.mmd** dùng Mermaid cho kiến trúc, trình tự, pipeline và trạng thái.',
      '- Hình 3, 5 và 6 có nhiều chi tiết; bố trí trang riêng, ưu tiên trang khổ ngang nếu hình rộng. Điều chỉnh số hình theo chương trong báo cáo.',
      '- Màu chỉ hỗ trợ phân nhóm; không mã hóa thêm trạng thái nghiệp vụ. Mũi tên nét đứt trong sequence là phản hồi; mũi tên submit bất đồng bộ chuyển giao thực thi cho executor.', '',
      '## Danh mục và vị trí đề xuất', '',
      md_table(['Hình', 'Tên hình', 'Vị trí trong báo cáo'], [[f['figure_number'], f['title'], f['section']] for f in figures]), '',
      '## Quy ước tên gọi', '', md_table(['Thuật ngữ', 'Nghĩa sử dụng'], content['terms']), '',
      '## Các sơ đồ và thuyết minh', '']
for f in figures:
    n = f['figure_number']
    ext = 'puml' if f['id'] in ('01-ca-su-dung', '03-hoat-dong-phan-tich') else 'mmd'
    md.extend([f'### Hình {n}. {f["title"]}', '',
               f'**Loại:** {f["type"]}. **Vị trí đề xuất:** {f["section"]}.', '',
               f'![Hình {n}. {f["title"]}](hinh/svg/{f["id"]}.svg)', '',
               f'**Chú thích dùng trong báo cáo:** {f["caption"]}', '',
               f'**Thuyết minh:** {f["description"]}', '',
               '**Đối chiếu code:** ' + ', '.join(source_link(s) for s in f['sources']) + '.', '',
               f'[SVG](hinh/svg/{f["id"]}.svg) · [PNG](hinh/png/{f["id"]}.png) · [PDF](hinh/pdf/{f["id"]}.pdf) · [Nguồn chỉnh sửa](nguon/{f["id"]}.{ext})', ''])
md.extend(['## Đối chiếu API', '', md_table(['Method', 'Endpoint', 'Chức năng', 'HTTP status chính'], content['api']), '',
           'Các mã lỗi trong bảng là những nhánh chính của luồng mô tả; lỗi hạ tầng hoặc file không tồn tại có thể phát sinh thêm. Mã 422 cũng do FastAPI/Pydantic tạo khi cấu trúc request không hợp lệ.', '',
           '## Danh mục 19 lớp đầu ra', '',
           'Tên tiếng Anh là giá trị nguyên gốc trong CLASS_NAMES và được trả qua API. Tên tiếng Việt dưới đây là cách gọi đề xuất trong phần thuyết minh; không thay đổi dữ liệu hoặc nhãn model.', '',
           md_table(['Mã lớp', 'Tên nguyên gốc trong code', 'Tên tiếng Việt đề xuất'], class_rows), '',
           '## Những quy tắc cần giữ khi diễn giải', ''])
md.extend(f'{i}. {rule}' for i, rule in enumerate(content['rules'], 1))
md.extend(['', '## Xuất lại tài liệu', '',
           'Công cụ tạo hình được cài riêng với ứng dụng. Script không chạy inference hoặc chỉnh sửa source của web.', '',
           '```bash',
           'PUPPETEER_SKIP_DOWNLOAD=true npm install --prefix /tmp/crop-map-report-tools @mermaid-js/mermaid-cli',
           '# Đặt plantuml.jar trong /tmp/crop-map-report-tools/; bộ hiện tại dùng PlantUML 1.2025.10.',
           'python3 -m pip install --target /tmp/crop-map-report-tools/python python-docx pypdf',
           'node docs/bao-cao/xuat-hinh.mjs',
           'PYTHONPATH=/tmp/crop-map-report-tools/python python3 docs/bao-cao/dong-goi.py',
           '```', '',
           'Có thể đặt REPORT_TOOLS_DIR và REPORT_CHROME khi thay đổi đường dẫn công cụ. Nếu Chrome nằm ở vị trí khác, cập nhật cả nguon/puppeteer-config.json dùng bởi Mermaid CLI. Tệp kiem-chung.json ghi hash các file nguồn được đối chiếu và kiểm tra sản phẩm xuất. Các tệp trong gói ZIP không chứa ảnh vệ tinh đầu vào, checkpoint hay thông tin khóa truy cập.', ''])
(HERE / 'README.md').write_text('\n'.join(md))

writer = PdfWriter()
verification = {'date': content['date'], 'scope': content['scope'], 'source_hashes': {}, 'figures': []}
for f in figures:
    reader = PdfReader(HERE / 'hinh/pdf' / (f['id'] + '.pdf'))
    if len(reader.pages) != 1:
        raise RuntimeError(f'{f["id"]}: expected one PDF page, got {len(reader.pages)}')
    writer.append(reader)
    with Image.open(HERE / 'hinh/png' / (f['id'] + '.png')) as img:
        img.verify()
    verification['figures'].append({'id': f['id'], 'pdf_pages': 1, 'svg': True, 'png_verified': True})
with (HERE / 'So-do-Crop-Map.pdf').open('wb') as handle:
    writer.write(handle)
assert len(PdfReader(HERE / 'So-do-Crop-Map.pdf').pages) == len(figures)

doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(7)
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(1.8)
doc.add_heading(content['title'], 0)
doc.add_paragraph('Ngày đối chiếu mã nguồn: ' + content['date'])
doc.add_paragraph(content['scope'])
doc.add_paragraph('Tài liệu gồm 12 hình có tên và thuật ngữ thống nhất. Mỗi hình có chú thích dùng trong báo cáo, thuyết minh và vị trí mã nguồn. Hình được đặt trên trang A3/A2 để giữ khả năng đọc; có thể sao chép SVG/PNG và chú thích vào báo cáo theo bố cục của trường.')


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = str(value)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    return t


doc.add_heading('Danh mục hình', 1)
table(['Hình', 'Tên', 'Mục báo cáo'], [[f['figure_number'], f['title'], f['section']] for f in figures])
doc.add_heading('Quy ước tên gọi', 1)
table(['Thuật ngữ', 'Ý nghĩa'], content['terms'])
for f in figures:
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    short, long = (42, 59.4) if f['pdf_paper'] == 'A2' else (29.7, 42)
    page_w, page_h = (short, long) if f['portrait'] else (long, short)
    section.page_width, section.page_height = Cm(page_w), Cm(page_h)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.5)
    doc.add_heading(f'Hình {f["figure_number"]}. {f["title"]}', 1)
    doc.add_paragraph(f'{f["type"]} — {f["section"]}')
    # Reserve room for caption, explanation and references under the drawing.
    available_w = page_w - 3
    available_h = page_h - 12
    ratio = min(available_w / f['svg_width'], available_h / f['svg_height'])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(HERE / 'hinh/png' / (f['id'] + '.png')), width=Cm(f['svg_width'] * ratio))
    p = doc.add_paragraph(f'Hình {f["figure_number"]}. {f["caption"]}')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
    doc.add_paragraph(f['description'])
    p = doc.add_paragraph('Đối chiếu: ' + ', '.join(f['sources']))
    for run in p.runs:
        run.font.size = Pt(10)

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
section.page_width, section.page_height = Cm(21), Cm(29.7)
doc.add_heading('Đối chiếu API', 1)
table(['Method', 'Endpoint', 'Chức năng', 'HTTP status chính'], content['api'])
doc.add_heading('Danh mục lớp đầu ra', 1)
doc.add_paragraph('Tên tiếng Anh là nhãn nguyên gốc trong code. Tên tiếng Việt là cách gọi đề xuất cho phần thuyết minh.')
table(['Mã', 'Tên trong code', 'Tên tiếng Việt'], class_rows)
doc.add_heading('Quy tắc diễn giải', 1)
for rule in content['rules']:
    doc.add_paragraph(rule, style='List Bullet')
doc.save(HERE / 'So-do-Crop-Map.docx')

for name in ['app.py', 'inference.py', 'imagery.py', 'model.py', 'web/app.js', 'compose.yaml', 'Dockerfile', 'weights/config.yaml']:
    verification['source_hashes'][name] = hashlib.sha256((PROJECT / name).read_bytes()).hexdigest()
verification['total_figures'] = len(figures)
verification['combined_pdf_pages'] = len(PdfReader(HERE / 'So-do-Crop-Map.pdf').pages)
verification['docx_zip_verified'] = zipfile.ZipFile(HERE / 'So-do-Crop-Map.docx').testzip() is None
(HERE / 'kiem-chung.json').write_text(json.dumps(verification, ensure_ascii=False, indent=2) + '\n')

archive = HERE / 'Bo-so-do-Crop-Map.zip'
with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as output:
    for p in sorted(HERE.rglob('*')):
        if p.is_file() and p != archive and '__pycache__' not in p.parts:
            output.write(p, Path('Crop-Map-So-Do') / p.relative_to(HERE))
print(json.dumps({'figures': len(figures), 'pdf_pages': verification['combined_pdf_pages'], 'docx': str(HERE / 'So-do-Crop-Map.docx'), 'zip': str(archive), 'zip_mb': round(archive.stat().st_size / 1024**2, 2)}, ensure_ascii=False))
